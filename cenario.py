"""Gera um arquivo de cenário a partir do modelo e da TIMETABLE nova."""

from __future__ import annotations

from copy import copy
from datetime import datetime, time
from io import BytesIO
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from compare import _cell_str, _normalize_column_name

MODELO_CENARIO_PADRAO = Path(__file__).parent / "Modelo Cenario.xlsx"

HEADER_ROW = 2
DATA_START_ROW = 3

PREFERRED_SOURCES: dict[str, list[str]] = {
    "linha": ["codigo linha", "código linha", "linha"],
    "frota": ["frota real", "frota", "frota teste", "classe"],
    "partida": ["1º embarque", "1o embarque", "primeiro embarque", "partida"],
    "chegada": [
        "último desembarque",
        "ultimo desembarque",
        "chegada",
    ],
    "local (lib.)": ["local (lib.)", "local (lib)"],
    "local (rec.)": ["local (rec.)", "local (rec)"],
    "d.o. (lib.)": ["d.o. (lib.)", "d.o. (lib)"],
    "d.o. (rec.)": ["d.o. (rec.)", "d.o. (rec)"],
    "serviço": ["serviço", "servico"],
    "t. e.": ["t. e.", "t.e.", "te"],
    "t.d.": ["t.d.", "td"],
}

TP_GR = time(1, 30)
TP_OUTROS = time(0, 45)

UM_MINUTO = time(0, 1)


def _norm(name: Any) -> str:
    return _normalize_column_name(name)


def _is_etapa(name: str) -> bool:
    n = _norm(name)
    return n == "etapa" or n.startswith("etapa/")


def _is_tp(name: str) -> bool:
    n = _norm(name).replace(" ", "")
    return n in {"t.p.", "t.p", "tp"}


def _is_te(name: str) -> bool:
    n = _norm(name).replace(" ", "")
    return n in {"t.e.", "t.e", "te"}


def _is_td(name: str) -> bool:
    n = _norm(name).replace(" ", "")
    return n in {"t.d.", "t.d", "td"}


def _is_tipo(name: str) -> bool:
    return _norm(name) == "tipo"


def _is_servico(name: str) -> bool:
    return _norm(name) in {"serviço", "servico"}


def _is_local_rec(name: str) -> bool:
    n = _norm(name).replace(" ", "")
    return n.startswith("local(rec")


def _source_lookup(df: pd.DataFrame) -> dict[str, str]:
    return {_norm(col): col for col in df.columns}


def resolve_source_column(scenario_col: str, lookup: dict[str, str]) -> str | None:
    n = _norm(scenario_col)
    preferred = PREFERRED_SOURCES.get(n)
    if preferred:
        for alias in preferred:
            if alias in lookup:
                return lookup[alias]
    if n in lookup:
        return lookup[n]
    compact = n.replace(" ", "")
    for key, original in lookup.items():
        if key.replace(" ", "") == compact:
            return original
    return None


def _parse_time(value: Any) -> time | None:
    if value is None or (isinstance(value, float) and pd.isna(value)):
        return None
    if isinstance(value, time):
        return value
    if isinstance(value, datetime):
        return value.time()
    text = _cell_str(value)
    if not text:
        return None
    for fmt in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    return None


def _parse_value(value: Any, sample: Any = None) -> Any:
    if isinstance(sample, time) or (
        hasattr(sample, "strftime") and not isinstance(sample, datetime)
    ):
        parsed = _parse_time(value)
        if parsed is not None:
            return parsed
    parsed_time = _parse_time(value)
    if parsed_time is not None and isinstance(sample, time):
        return parsed_time
    text = _cell_str(value)
    if text == "":
        return None
    if isinstance(sample, (int, float)) and not isinstance(sample, bool):
        try:
            number = float(text.replace(",", "."))
            if number.is_integer():
                return int(number)
            return number
        except ValueError:
            return text
    if text.lstrip("-").isdigit():
        try:
            return int(text)
        except ValueError:
            return text
    if parsed_time is not None:
        return parsed_time
    return text


def _copy_style(source, target) -> None:
    if source.has_style:
        target.font = copy(source.font)
        target.border = copy(source.border)
        target.fill = copy(source.fill)
        target.alignment = copy(source.alignment)
        target.protection = copy(source.protection)
    target.number_format = source.number_format


def _headers(ws: Worksheet, header_row: int = HEADER_ROW) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for col in range(1, (ws.max_column or 1) + 1):
        value = ws.cell(header_row, col).value
        if value is None or str(value).strip() == "":
            continue
        mapping[str(value).strip()] = col
    return mapping


def _local_rec_starts_with_gr(value: Any) -> bool:
    text = _cell_str(value).strip().casefold()
    return text.startswith("gr")


def tp_for_local_rec(local_rec: Any) -> time:
    return TP_GR if _local_rec_starts_with_gr(local_rec) else TP_OUTROS


def build_cenario_excel(
    timetable: pd.DataFrame,
    modelo_path: str | Path | None = None,
    frota: pd.DataFrame | None = None,
    timetable_old: pd.DataFrame | None = None,
) -> bytes:
    path = Path(modelo_path) if modelo_path else MODELO_CENARIO_PADRAO
    if not path.exists():
        raise FileNotFoundError(
            f"Modelo de cenário não encontrado: {path}"
        )

    wb = load_workbook(path)
    ws = wb.active
    headers = _headers(ws)
    if not headers:
        raise ValueError("O modelo de cenário não tem cabeçalho na linha 2.")

    lookup = _source_lookup(timetable)

    # Aba FROTA do arquivo novo (antiga Planilha3):
    #   coluna A (sem cabeçalho) = Prefixo
    #   coluna K = Frota
    #
    # O mapa é usado para preencher a coluna FROTA do Novo Cenário
    # a partir do Prefixo existente na TIMETABLE.
    frota_por_prefixo: dict[str, Any] = {}
    if frota is not None and not frota.empty:
        if frota.shape[1] < 11:
            raise ValueError(
                "A aba FROTA precisa ter pelo menos 11 colunas "
                "(A = Prefixo e K = Frota)."
            )

        for _, row in frota.iterrows():
            prefixo = row.iloc[0]
            frota_valor = row.iloc[10]

            prefixo_key = _cell_str(prefixo).strip()
            if prefixo_key:
                frota_por_prefixo[prefixo_key] = frota_valor

    # Localiza a coluna Prefixo na TIMETABLE.
    prefixo_source_col = None
    for col in timetable.columns:
        if _norm(col) == "prefixo":
            prefixo_source_col = col
            break

    sample_row = DATA_START_ROW
    last_template_row = ws.max_row or DATA_START_ROW
    default_values = {
        col_idx: ws.cell(sample_row, col_idx).value
        for col_idx in headers.values()
    }

    local_rec_col_name = next(
        (name for name in headers if _is_local_rec(name)),
        None,
    )
    etapa_cols = [name for name in headers if _is_etapa(name)]
    tp_cols = [name for name in headers if _is_tp(name)]

    tipo_col_name = next((name for name in headers if _is_tipo(name)), None)

    # Localiza FROTA no Novo Cenário.
    frota_col_name = next(
        (name for name in headers if _norm(name) == "frota"),
        None,
    )

    # Localiza SERVIÇO no Novo Cenário e a coluna correspondente
    # na TIMETABLE antiga (mesmo nome/alias, resolvido via
    # PREFERRED_SOURCES/resolve_source_column).
    servico_col_name = next(
        (name for name in headers if _is_servico(name)),
        None,
    )

    old_servico_source_col = None
    if servico_col_name is not None and timetable_old is not None:
        old_lookup = _source_lookup(timetable_old)
        old_servico_source_col = resolve_source_column(
            servico_col_name,
            old_lookup,
        )

    n_rows = len(timetable)
    target_last = DATA_START_ROW + n_rows - 1 if n_rows else DATA_START_ROW - 1

    if last_template_row > max(target_last, DATA_START_ROW - 1):
        ws.delete_rows(
            target_last + 1,
            last_template_row - target_last,
        )

    for offset in range(n_rows):
        excel_row = DATA_START_ROW + offset
        src = timetable.iloc[offset]
        local_rec_value = None

        # Busca o Prefixo da TIMETABLE na coluna A da aba FROTA
        # e usa o valor da coluna K como FROTA.
        frota_por_prefixo_valor = None
        if frota_col_name and prefixo_source_col is not None:
            prefixo_key = _cell_str(src[prefixo_source_col]).strip()
            if prefixo_key in frota_por_prefixo:
                frota_por_prefixo_valor = frota_por_prefixo[prefixo_key]

        # Serviço vem da TIMETABLE antiga, na MESMA posição (linha)
        # da TIMETABLE nova. Se não houver linha correspondente
        # (ex.: linha nova adicionada, sem equivalente na antiga),
        # cai no comportamento padrão (TIMETABLE nova) mais abaixo.
        servico_old_value = None
        if (
            old_servico_source_col is not None
            and timetable_old is not None
            and offset < len(timetable_old)
        ):
            servico_old_value = timetable_old.iloc[offset][old_servico_source_col]

        if local_rec_col_name:
            source_col = resolve_source_column(local_rec_col_name, lookup)
            if source_col is not None:
                local_rec_value = src[source_col]

        for name, col_idx in headers.items():
            sample_cell = ws.cell(sample_row, col_idx)
            cell = ws.cell(excel_row, col_idx)
            if excel_row != sample_row:
                _copy_style(sample_cell, cell)

            if _is_etapa(name):
                cell.value = 1
                continue

            if _is_tp(name):
                cell.value = tp_for_local_rec(local_rec_value)
                continue

            if _is_tipo(name):
                cell.value = "FTR"
                continue

            # T.E. e T.D. sempre 1 minuto.
            if _is_te(name) or _is_td(name):
                cell.value = UM_MINUTO
                continue

            # FROTA vem exclusivamente do cruzamento:
            # TIMETABLE[Prefixo] -> FROTA[A] -> FROTA[K].
            if frota_col_name and name == frota_col_name:
                cell.value = _parse_value(
                    frota_por_prefixo_valor,
                    sample_cell.value,
                )
                continue

            # SERVIÇO vem da TIMETABLE antiga, por posição.
            if servico_col_name and name == servico_col_name:
                if servico_old_value is not None:
                    cell.value = _parse_value(servico_old_value, sample_cell.value)
                else:
                    # Sem linha correspondente na antiga: usa a
                    # TIMETABLE nova como alternativa.
                    fallback_col = resolve_source_column(name, lookup)
                    cell.value = (
                        _parse_value(src[fallback_col], sample_cell.value)
                        if fallback_col is not None
                        else default_values.get(col_idx)
                    )
                continue

            source_col = resolve_source_column(name, lookup)
            if source_col is None:
                cell.value = default_values.get(col_idx)
                continue

            cell.value = _parse_value(src[source_col], sample_cell.value)

        for name in etapa_cols:
            ws.cell(excel_row, headers[name]).value = 1
        for name in tp_cols:
            ws.cell(excel_row, headers[name]).value = tp_for_local_rec(
                local_rec_value
            )

    out = BytesIO()
    wb.save(out)
    return out.getvalue()