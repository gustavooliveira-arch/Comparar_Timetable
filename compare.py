"""
Compara a aba TIMETABLE de dois arquivos Excel.

Tipos de alterações de células:
- alterada: valor antigo != valor novo, ambos preenchidos
- adicionada: vazio -> valor
- removida: valor -> vazio

Também identifica:
- linhas adicionadas
- linhas removidas
- colunas exclusivas do arquivo antigo
- colunas exclusivas do arquivo novo
- colunas ocultas/visíveis
"""

from __future__ import annotations

from dataclasses import dataclass, field
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SHEET_NAME = "TIMETABLE"


# ============================================================
# UTILITÁRIOS
# ============================================================

def _cell_str(value: Any) -> str:
    """Converte o conteúdo de uma célula para texto comparável."""

    if value is None:
        return ""

    try:
        if pd.isna(value):
            return ""
    except Exception:
        pass

    if isinstance(value, pd.Timestamp):
        if value.hour or value.minute or value.second:
            return value.strftime("%Y-%m-%d %H:%M:%S")
        return value.strftime("%Y-%m-%d")

    text = str(value).strip()

    # Converte "1.0" para "1"
    if (
        text.endswith(".0")
        and text.replace(".", "", 1).replace("-", "", 1).isdigit()
    ):
        try:
            as_float = float(text)
            if as_float.is_integer():
                return str(int(as_float))
        except ValueError:
            pass

    return text


def _normalize_column_name(name: Any) -> str:
    """
    Normaliza o nome da coluna para comparação.

    Exemplo:
        " Teste " -> "teste"
        "TESTE"   -> "teste"
        "Teste  1" -> "teste 1"
    """

    text = str(name).strip()
    text = " ".join(text.split())
    return text.casefold()


# ============================================================
# COLUNAS
# ============================================================

@dataclass
class ColumnInfo:
    name: str
    hidden: bool = False

    @property
    def visibility(self) -> str:
        return "Oculta" if self.hidden else "Visível"


def _is_ghost_name(name: Any) -> bool:
    text = str(name).strip()
    lower = text.lower()

    return (
        text == ""
        or lower in {"nan", "none", "unnamed"}
        or lower.startswith("unnamed:")
        or lower.startswith("col_")
    )


def _drop_ghost_columns(df: pd.DataFrame) -> pd.DataFrame:
    keep = []
    for col in df.columns:
        if _is_ghost_name(col):
            continue
        keep.append(col)
    if not keep:
        return df.iloc[:, 0:0]
    return df[keep]


def _unique_headers(names: list[Any]) -> list[str]:
    """
    Garante que nomes de colunas duplicados sejam únicos.

    Exemplo:
        Teste
        Teste
        Teste

    vira:

        Teste
        Teste_1
        Teste_2
    """

    result: list[str] = []
    counters: dict[str, int] = {}

    for raw_name in names:
        name = str(raw_name).strip()

        if name == "":
            name = "Coluna"

        if name not in counters:
            counters[name] = 0
            result.append(name)
        else:
            counters[name] += 1
            result.append(f"{name}_{counters[name]}")

    return result


# ============================================================
# CABEÇALHO / LEITURA DO EXCEL
# ============================================================

def detect_header_row(
    raw: pd.DataFrame,
    max_scan: int = 20,
) -> int:
    """
    Detecta a linha do cabeçalho.

    Retorna índice zero-based.
    """

    if raw.empty:
        return 0

    scan = min(max_scan, len(raw))

    best_row = 0
    best_score = -1

    for row_idx in range(scan):
        row = raw.iloc[row_idx]

        values = [
            _cell_str(value)
            for value in row.tolist()
            if _cell_str(value) != ""
        ]

        if not values:
            continue

        # Quanto mais células preenchidas, maior a chance de ser cabeçalho.
        unique_values = len(set(values))
        score = len(values) * 10 + unique_values

        if score > best_score:
            best_score = score
            best_row = row_idx

    return best_row


def normalize_frame(df: pd.DataFrame) -> pd.DataFrame:
    """Normaliza um DataFrame já carregado (testes e comparações manuais)."""
    out = df.copy()
    out.columns = _unique_headers(list(out.columns))
    out = out.dropna(how="all")
    for col in out.columns:
        out[col] = out[col].map(_cell_str)
    return _drop_ghost_columns(out).reset_index(drop=True)


def _frame_from_raw(
    raw: pd.DataFrame,
    header_idx: int,
) -> pd.DataFrame:

    if raw.empty:
        return pd.DataFrame()

    headers = _unique_headers(
        raw.iloc[header_idx].tolist()
    )

    data = raw.iloc[header_idx + 1:].copy()
    data.columns = headers
    data = data.dropna(how="all")

    for col in data.columns:
        data[col] = data[col].map(_cell_str)

    data = _drop_ghost_columns(data)

    return data.reset_index(drop=True)


def get_column_visibility(
    path_or_buffer,
    sheet_name: str,
) -> dict[str, bool]:
    """
    Retorna:
        {
            "Teste": False,
            "Coluna X": True
        }
    """

    try:
        if hasattr(path_or_buffer, "seek"):
            path_or_buffer.seek(0)

        wb = load_workbook(
            path_or_buffer,
            read_only=True,
            data_only=True,
            keep_links=False,
        )

        if sheet_name not in wb.sheetnames:
            wb.close()
            return {}

        ws = wb[sheet_name]

        max_col = ws.max_column or 1
        max_scan = min(ws.max_row or 1, 20)

        raw_rows = []

        for row in range(1, max_scan + 1):
            raw_rows.append(
                [
                    ws.cell(row, col).value
                    for col in range(1, max_col + 1)
                ]
            )

        raw = pd.DataFrame(raw_rows)

        header_idx = detect_header_row(raw)
        header_excel_row = header_idx + 1

        visibility: dict[str, bool] = {}

        original_headers = _unique_headers(
            raw.iloc[header_idx].tolist()
        )

        for col_idx, name in enumerate(
            original_headers,
            start=1,
        ):
            if _is_ghost_name(name):
                continue

            letter = get_column_letter(col_idx)

            visibility[name] = (
                ws.column_dimensions[letter].hidden is True
            )

        wb.close()
        return visibility

    except Exception:
        return {}


def _read_sheet_raw(path_or_buffer, sheet_name: str) -> tuple[pd.DataFrame, str]:
    last_error: Exception | None = None
    sheets = None

    for engine in ("calamine", None):
        try:
            if hasattr(path_or_buffer, "seek"):
                path_or_buffer.seek(0)
            kwargs: dict[str, Any] = {"sheet_name": None, "header": None}
            if engine:
                kwargs["engine"] = engine
            sheets = pd.read_excel(path_or_buffer, **kwargs)
            if isinstance(sheets, dict):
                break
        except Exception as exc:
            last_error = exc
            sheets = None

    if not isinstance(sheets, dict):
        raise ValueError(
            f"Não foi possível ler a aba '{sheet_name}': {last_error}"
        ) from last_error

    if sheet_name not in sheets:
        normalized = {
            str(name).strip().casefold(): name
            for name in sheets
        }
        found = normalized.get(str(sheet_name).strip().casefold())
        if found is None:
            raise ValueError(
                f"Aba '{sheet_name}' não encontrada. "
                f"Abas disponíveis: {', '.join(str(n) for n in sheets)}"
            )
        sheet_name = found

    return sheets[sheet_name], sheet_name


def load_timetable(
    path_or_buffer,
    sheet_name: str = SHEET_NAME,
    include_visibility: bool = False,
) -> pd.DataFrame:
    """Carrega a aba TIMETABLE."""

    try:
        raw, sheet_name = _read_sheet_raw(path_or_buffer, sheet_name)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(
            f"Não foi possível ler a aba '{sheet_name}': {exc}"
        ) from exc

    header_idx = detect_header_row(raw)
    df = _frame_from_raw(raw, header_idx)
    df.attrs["header_row"] = header_idx + 1
    df.attrs["column_visibility"] = {}

    if include_visibility:
        try:
            if hasattr(path_or_buffer, "seek"):
                path_or_buffer.seek(0)
            df.attrs["column_visibility"] = get_column_visibility(
                path_or_buffer,
                sheet_name,
            )
        except Exception:
            df.attrs["column_visibility"] = {}

    return df


# ============================================================
# ALTERAÇÃO DE CÉLULA
# ============================================================

@dataclass
class CellChange:
    """
    Representa uma alteração de célula.

    row:
        posição da linha no DataFrame, começando em 0.

    column:
        nome da coluna.

    old:
        valor antigo.

    new:
        valor novo.

    change_type:
        "alterada"
        "adicionada"
        "removida"
    """

    row: int
    column: str
    old: str
    new: str
    change_type: str = "alterada"


# ============================================================
# RESULTADO
# ============================================================

@dataclass
class CompareResult:

    added_rows: pd.DataFrame
    removed_rows: pd.DataFrame

    modified: list[CellChange] = field(
        default_factory=list
    )

    added_cells: list[CellChange] = field(
        default_factory=list
    )

    removed_cells: list[CellChange] = field(
        default_factory=list
    )

    extra_columns_old: list[ColumnInfo] = field(
        default_factory=list
    )

    extra_columns_new: list[ColumnInfo] = field(
        default_factory=list
    )

    @property
    def all_cell_changes(self) -> list[CellChange]:
        return (
            self.modified
            + self.added_cells
            + self.removed_cells
        )

    @property
    def has_changes(self) -> bool:
        return bool(
            self.modified
            or self.added_cells
            or self.removed_cells
            or not self.added_rows.empty
            or not self.removed_rows.empty
            or self.extra_columns_old
            or self.extra_columns_new
        )

    def modified_frame(self) -> pd.DataFrame:

        rows = []

        for change in self.modified:
            rows.append(
                {
                    "linha": change.row + 1,
                    "coluna": change.column,
                    "tipo": "Alterada",
                    "valor antigo": change.old,
                    "valor novo": change.new,
                }
            )

        for change in self.added_cells:
            rows.append(
                {
                    "linha": change.row + 1,
                    "coluna": change.column,
                    "tipo": "Adicionada",
                    "valor antigo": change.old,
                    "valor novo": change.new,
                }
            )

        for change in self.removed_cells:
            rows.append(
                {
                    "linha": change.row + 1,
                    "coluna": change.column,
                    "tipo": "Removida",
                    "valor antigo": change.old,
                    "valor novo": change.new,
                }
            )

        return pd.DataFrame(rows)


# ============================================================
# COLUNAS EXCLUSIVAS
# ============================================================

def _column_info(
    df: pd.DataFrame,
    column_name: str,
) -> ColumnInfo:

    visibility = df.attrs.get(
        "column_visibility",
        {},
    )

    hidden = bool(
        visibility.get(column_name, False)
    )

    return ColumnInfo(
        name=column_name,
        hidden=hidden,
    )


def _extra_columns(
    old: pd.DataFrame,
    new: pd.DataFrame,
) -> tuple[list[ColumnInfo], list[ColumnInfo]]:

    old_map = {
        _normalize_column_name(c): c
        for c in old.columns
    }

    new_map = {
        _normalize_column_name(c): c
        for c in new.columns
    }

    old_only = []

    for normalized, original in old_map.items():
        if normalized not in new_map:
            old_only.append(
                _column_info(old, original)
            )

    new_only = []

    for normalized, original in new_map.items():
        if normalized not in old_map:
            new_only.append(
                _column_info(new, original)
            )

    return old_only, new_only


# ============================================================
# COMPARAÇÃO POR POSIÇÃO
# ============================================================

def compare_positional(
    old: pd.DataFrame,
    new: pd.DataFrame,
) -> CompareResult:

    old = old.copy()
    new = new.copy()

    old_only_cols, new_only_cols = _extra_columns(
        old,
        new,
    )

    old_map = {
        _normalize_column_name(c): c
        for c in old.columns
    }

    new_map = {
        _normalize_column_name(c): c
        for c in new.columns
    }

    shared_normalized = [
        key
        for key in old_map
        if key in new_map
    ]

    shared = [
        (
            old_map[key],
            new_map[key],
        )
        for key in shared_normalized
    ]

    # --------------------------------------------------------
    # LINHAS
    # --------------------------------------------------------

    min_rows = min(
        len(old),
        len(new),
    )

    added_rows = pd.DataFrame()

    if len(new) > len(old):
        added_rows = new.iloc[len(old):].copy()

        added_rows.insert(
            0,
            "_linha_excel",
            range(
                len(old) + 1,
                len(new) + 1,
            ),
        )

    removed_rows = pd.DataFrame()

    if len(old) > len(new):
        removed_rows = old.iloc[len(new):].copy()

        removed_rows.insert(
            0,
            "_linha_excel",
            range(
                len(new) + 1,
                len(old) + 1,
            ),
        )

    modified: list[CellChange] = []
    added_cells: list[CellChange] = []
    removed_cells: list[CellChange] = []

    # --------------------------------------------------------
    # COLUNAS EM COMUM (vetorizado)
    #
    # Em vez de percorrer célula a célula com duas iterações Python,
    # cada coluna é convertida para texto uma única vez (`.map`) e a
    # comparação vira uma operação de array NumPy. O loop restante só
    # toca nas células efetivamente diferentes.
    # --------------------------------------------------------

    if shared and min_rows > 0:

        old_matrix = np.column_stack(
            [
                old[old_col]
                .iloc[:min_rows]
                .map(_cell_str)
                .to_numpy(object)
                for old_col, _ in shared
            ]
        )

        new_matrix = np.column_stack(
            [
                new[new_col]
                .iloc[:min_rows]
                .map(_cell_str)
                .to_numpy(object)
                for _, new_col in shared
            ]
        )

        for row_idx, col_pos in np.argwhere(
            old_matrix != new_matrix
        ):

            old_col, new_col = shared[col_pos]

            old_value = old_matrix[row_idx, col_pos]
            new_value = new_matrix[row_idx, col_pos]

            row_idx = int(row_idx)

            if old_value == "" and new_value != "":
                added_cells.append(
                    CellChange(
                        row=row_idx,
                        column=new_col,
                        old=old_value,
                        new=new_value,
                        change_type="adicionada",
                    )
                )

            elif old_value != "" and new_value == "":
                removed_cells.append(
                    CellChange(
                        row=row_idx,
                        column=old_col,
                        old=old_value,
                        new=new_value,
                        change_type="removida",
                    )
                )

            else:
                modified.append(
                    CellChange(
                        row=row_idx,
                        column=new_col,
                        old=old_value,
                        new=new_value,
                        change_type="alterada",
                    )
                )

    # --------------------------------------------------------
    # COLUNAS NOVAS (vetorizado)
    # --------------------------------------------------------

    for key in new_map:

        if key in old_map:
            continue

        new_col = new_map[key]

        col_values = (
            new[new_col]
            .iloc[:min_rows]
            .map(_cell_str)
            .to_numpy(object)
        )

        for row_idx in np.argwhere(
            col_values != ""
        ).flatten():

            added_cells.append(
                CellChange(
                    row=int(row_idx),
                    column=new_col,
                    old="",
                    new=col_values[row_idx],
                    change_type="adicionada",
                )
            )

    # --------------------------------------------------------
    # COLUNAS REMOVIDAS (vetorizado)
    # --------------------------------------------------------

    for key in old_map:

        if key in new_map:
            continue

        old_col = old_map[key]

        col_values = (
            old[old_col]
            .iloc[:min_rows]
            .map(_cell_str)
            .to_numpy(object)
        )

        for row_idx in np.argwhere(
            col_values != ""
        ).flatten():

            removed_cells.append(
                CellChange(
                    row=int(row_idx),
                    column=old_col,
                    old=col_values[row_idx],
                    new="",
                    change_type="removida",
                )
            )

    return CompareResult(
        added_rows=added_rows.reset_index(drop=True),
        removed_rows=removed_rows.reset_index(drop=True),
        modified=modified,
        added_cells=added_cells,
        removed_cells=removed_cells,
        extra_columns_old=old_only_cols,
        extra_columns_new=new_only_cols,
    )
