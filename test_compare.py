from io import BytesIO

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.workbook import Workbook

from compare import compare_by_keys, compare_positional, load_timetable, normalize_frame
from export import build_updated_excel


def _xlsx(df: pd.DataFrame, sheet: str = "TIMETABLE") -> BytesIO:
    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name=sheet, index=False)
    buf.seek(0)
    return buf


def test_positional_detects_cell_change_and_new_row():
    old = normalize_frame(
        pd.DataFrame(
            {
                "Dia": ["Seg", "Ter"],
                "Aula": ["Math", "Hist"],
                "Sala": ["A1", "B2"],
            }
        )
    )
    new = normalize_frame(
        pd.DataFrame(
            {
                "Dia": ["Seg", "Ter", "Qua"],
                "Aula": ["Math", "Geo", "Bio"],
                "Sala": ["A1", "B2", "C3"],
            }
        )
    )
    result = compare_positional(old, new)
    assert len(result.modified) == 1
    assert result.modified[0].column == "Aula"
    assert result.modified[0].old == "Hist"
    assert result.modified[0].new == "Geo"
    assert len(result.added_rows) == 1
    assert result.added_rows.iloc[0]["Dia"] == "Qua"


def test_by_key_ignores_row_order():
    old = normalize_frame(
        pd.DataFrame({"Codigo": ["T1", "T2"], "Sala": ["A", "B"]})
    )
    new = normalize_frame(
        pd.DataFrame({"Codigo": ["T2", "T1"], "Sala": ["B", "C"]})
    )
    result = compare_by_keys(old, new, ["Codigo"])
    assert result.added_rows.empty
    assert result.removed_rows.empty
    assert len(result.modified) == 1
    assert result.modified[0].old == "A"
    assert result.modified[0].new == "C"


def test_load_timetable_reads_named_sheet():
    buf = _xlsx(pd.DataFrame({"Dia": ["Seg"], "Aula": ["Math"]}))
    df = load_timetable(buf)
    assert list(df.columns) == ["Dia", "Aula"]
    assert df.iloc[0]["Aula"] == "Math"


def _styled_old_workbook() -> bytes:
    wb = Workbook()
    extra = wb.active
    extra.title = "CAPA"
    extra["A1"] = "Capa"
    extra["A1"].font = Font(bold=True, color="FF0000")
    ws = wb.create_sheet("TIMETABLE")
    ws["A1"] = "Dia"
    ws["B1"] = "Aula"
    ws["C1"] = "Sala"
    ws["A2"] = "Seg"
    ws["B2"] = "Math"
    ws["C2"] = "A1"
    ws["A3"] = "Ter"
    ws["B3"] = "Hist"
    ws["C3"] = "B2"
    ws.column_dimensions["B"].width = 22
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_export_keeps_other_sheets_and_applies_changes():
    old_bytes = _styled_old_workbook()
    old = normalize_frame(pd.DataFrame({"Dia": ["Seg", "Ter"], "Aula": ["Math", "Hist"], "Sala": ["A1", "B2"]}))
    new = normalize_frame(
        pd.DataFrame({"Dia": ["Seg", "Ter", "Qua"], "Aula": ["Math", "Geo", "Bio"], "Sala": ["A1", "B2", "C3"]})
    )
    result = compare_positional(old, new)
    out = build_updated_excel(old_bytes, result, "TIMETABLE")
    wb = load_workbook(BytesIO(out))
    assert wb.sheetnames == ["CAPA", "TIMETABLE"]
    assert wb["CAPA"]["A1"].value == "Capa"
    assert wb["CAPA"]["A1"].font.bold is True
    ws = wb["TIMETABLE"]
    assert ws["B3"].value == "Geo"
    assert ws["B3"].fill.fgColor.rgb in ("00FFF2CC", "FFF2CC")
    assert ws["A4"].value == "Qua"
    assert ws["B4"].value == "Bio"
    assert ws["B4"].fill.fgColor.rgb in ("00C6EFCE", "C6EFCE")
    assert ws.column_dimensions["B"].width == 22
