"""
Gera arquivos Excel a partir do resultado da comparação.

build_updated_excel:
    Mantém o arquivo antigo completo e aplica as alterações.

build_diff_only_excel:
    Gera um arquivo contendo somente as células envolvidas
    nas alterações.
"""

from __future__ import annotations

from copy import copy
from io import BytesIO
from typing import Any

import pandas as pd

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from compare import (
    CompareResult,
    _cell_str,
    _normalize_column_name,
    detect_header_row,
)


# ============================================================
# CORES
# ============================================================

# Alteração: valor -> outro valor
FILL_CHANGED = PatternFill(
    fill_type="solid",
    fgColor="FFF2CC",
)

# Adição: vazio -> valor
FILL_ADDED = PatternFill(
    fill_type="solid",
    fgColor="C6EFCE",
)

# Remoção: valor -> vazio
FILL_REMOVED = PatternFill(
    fill_type="solid",
    fgColor="FFC7CE",
)


# ============================================================
# CABEÇALHO
# ============================================================

def _header_row(ws: Worksheet) -> int:

    max_col = ws.max_column or 1
    max_scan = min(
        ws.max_row or 1,
        20,
    )

    rows = []

    for row in range(
        1,
        max_scan + 1,
    ):
        rows.append(
            [
                ws.cell(row, col).value
                for col in range(
                    1,
                    max_col + 1,
                )
            ]
        )

    if not rows:
        return 1

    raw = pd.DataFrame(rows)

    return detect_header_row(raw) + 1


def _unique_headers_excel(
    ws: Worksheet,
    header_row: int,
) -> list[str]:

    names = []

    counters: dict[str, int] = {}

    for col in range(
        1,
        (ws.max_column or 1) + 1,
    ):

        value = ws.cell(
            header_row,
            col,
        ).value

        name = (
            str(value).strip()
            if value is not None
            else ""
        )

        if not name:
            name = f"Coluna_{col}"

        if name not in counters:
            counters[name] = 0
            names.append(name)
        else:
            counters[name] += 1
            names.append(
                f"{name}_{counters[name]}"
            )

    return names


def _headers(
    ws: Worksheet,
    header_row: int | None = None,
) -> dict[str, int]:

    row = (
        header_row
        if header_row is not None
        else _header_row(ws)
    )

    names = _unique_headers_excel(
        ws,
        row,
    )

    mapping: dict[str, int] = {}

    for col, name in enumerate(
        names,
        start=1,
    ):

        if not name:
            continue

        normalized = _normalize_column_name(
            name
        )

        # Mantém a primeira ocorrência.
        if normalized not in mapping:
            mapping[normalized] = col

    return mapping


# ============================================================
# VALORES
# ============================================================

def _coerce_value(
    text: str,
    sample: Any = None,
) -> Any:

    if text == "":
        return None

    # Se a célula original era data/hora.
    if hasattr(sample, "strftime"):

        try:
            ts = pd.to_datetime(
                text,
                errors="coerce",
            )

            if pd.notna(ts):
                return ts.to_pydatetime()

        except Exception:
            pass

    # Inteiro.
    if text.lstrip("-").isdigit():

        try:
            return int(text)

        except ValueError:
            return text

    # Decimal.
    try:

        if "." in text:
            return float(text)

    except ValueError:
        pass

    return text


def _write_cell(
    ws: Worksheet,
    row: int,
    col: int,
    text: str,
    fill: PatternFill | None = None,
) -> None:

    cell = ws.cell(
        row=row,
        column=col,
    )

    cell.value = _coerce_value(
        text,
        cell.value,
    )

    if fill is not None:
        cell.fill = copy(fill)


# ============================================================
# ESTILO
# ============================================================

def _copy_cell_style(
    source,
    target,
) -> None:

    if not source.has_style:
        return

    target.font = copy(source.font)
    target.fill = copy(source.fill)
    target.border = copy(source.border)
    target.alignment = copy(source.alignment)
    target.number_format = source.number_format
    target.protection = copy(source.protection)


def _copy_header_style(
    ws: Worksheet,
    from_col: int,
    to_col: int,
    header_row: int = 1,
) -> None:

    source = ws.cell(
        header_row,
        from_col,
    )

    target = ws.cell(
        header_row,
        to_col,
    )

    _copy_cell_style(
        source,
        target,
    )

    source_letter = get_column_letter(
        from_col
    )

    target_letter = get_column_letter(
        to_col
    )

    source_dimension = (
        ws.column_dimensions[
            source_letter
        ]
    )

    target_dimension = (
        ws.column_dimensions[
            target_letter
        ]
    )

    if source_dimension.width:
        target_dimension.width = (
            source_dimension.width
        )

    target_dimension.hidden = (
        source_dimension.hidden is True
    )


# ============================================================
# COLUNAS
# ============================================================

def _ensure_column(
    ws: Worksheet,
    headers: dict[str, int],
    name: str,
    header_row: int = 1,
) -> int:

    normalized = _normalize_column_name(
        name
    )

    if normalized in headers:
        return headers[normalized]

    col = (
        ws.max_column or 0
    ) + 1

    if headers:

        reference_col = next(
            iter(headers.values())
        )

        _copy_header_style(
            ws,
            reference_col,
            col,
            header_row,
        )

    ws.cell(
        header_row,
        col,
    ).value = name

    headers[normalized] = col

    return col


# ============================================================
# ÚLTIMA LINHA
# ============================================================

def _last_data_row(
    ws: Worksheet,
    header_row: int = 1,
) -> int:

    max_row = ws.max_row or 1

    for row in range(
        max_row,
        header_row,
        -1,
    ):

        if any(
            ws.cell(row, col).value
            not in (None, "")
            for col in range(
                1,
                (ws.max_column or 1) + 1,
            )
        ):
            return row

    return header_row


# ============================================================
# TEMPLATE
# ============================================================

def _load_template(
    old_bytes: bytes,
    sheet_name: str,
):

    try:

        wb = load_workbook(
            BytesIO(old_bytes)
        )

    except Exception:

        wb = Workbook()

        ws = wb.active
        ws.title = sheet_name

        return wb

    if sheet_name not in wb.sheetnames:
        wb.create_sheet(sheet_name)

    return wb


# ============================================================
# EXCEL ATUALIZADO
# ============================================================

def build_updated_excel(
    old_bytes: bytes,
    result: CompareResult,
    sheet_name: str,
    key_cols: list[str] | None = None,
) -> bytes:

    """
    Copia o arquivo antigo por inteiro e aplica as diferenças.

    Cores:

    amarelo = célula alterada
    verde   = célula adicionada
    vermelho = célula removida

    Linhas novas são adicionadas ao final.
    """

    wb = _load_template(
        old_bytes,
        sheet_name,
    )

    ws = wb[sheet_name]

    header_row = _header_row(ws)

    headers = _headers(
        ws,
        header_row,
    )

    # --------------------------------------------------------
    # CÉLULAS ALTERADAS
    # --------------------------------------------------------

    for change in result.modified:

        # change.row é zero-based no DataFrame.
        target_row = (
            header_row
            + change.row
            + 1
        )

        col = _ensure_column(
            ws,
            headers,
            change.column,
            header_row,
        )

        _write_cell(
            ws,
            target_row,
            col,
            change.new,
            FILL_CHANGED,
        )

    # --------------------------------------------------------
    # CÉLULAS ADICIONADAS
    # --------------------------------------------------------

    for change in result.added_cells:

        target_row = (
            header_row
            + change.row
            + 1
        )

        col = _ensure_column(
            ws,
            headers,
            change.column,
            header_row,
        )

        _write_cell(
            ws,
            target_row,
            col,
            change.new,
            FILL_ADDED,
        )

    # --------------------------------------------------------
    # CÉLULAS REMOVIDAS
    # --------------------------------------------------------

    for change in result.removed_cells:

        target_row = (
            header_row
            + change.row
            + 1
        )

        col = _ensure_column(
            ws,
            headers,
            change.column,
            header_row,
        )

        # O valor novo é vazio.
        # Portanto limpamos a célula.
        _write_cell(
            ws,
            target_row,
            col,
            "",
            FILL_REMOVED,
        )

    # --------------------------------------------------------
    # LINHAS ADICIONADAS
    # --------------------------------------------------------

    added = result.added_rows.copy()

    added = added.drop(
        columns=["_linha_excel"],
        errors="ignore",
    )

    next_row = (
        _last_data_row(
            ws,
            header_row,
        )
        + 1
    )

    for _, row in added.iterrows():

        for col_name, value in row.items():

            col = _ensure_column(
                ws,
                headers,
                str(col_name),
                header_row,
            )

            _write_cell(
                ws,
                next_row,
                col,
                _cell_str(value),
                FILL_ADDED,
            )

        next_row += 1

    # --------------------------------------------------------
    # LINHAS REMOVIDAS
    # --------------------------------------------------------

    removed = result.removed_rows.copy()

    if "_linha_excel" in removed.columns:

        rows_to_delete = sorted(
            {
                int(value)
                for value in removed[
                    "_linha_excel"
                ]
            },
            reverse=True,
        )

        # _linha_excel é zero-based na comparação.
        # Converte para a posição real da planilha.
        for row_idx in rows_to_delete:

            excel_row = (
                header_row
                + row_idx
                + 1
            )

            if excel_row > header_row:
                ws.delete_rows(
                    excel_row
                )

    out = BytesIO()

    wb.save(out)

    return out.getvalue()


# ============================================================
# EXCEL SOMENTE ALTERAÇÕES
# ============================================================

def build_diff_only_excel(
    old_bytes: bytes,
    result: CompareResult,
    sheet_name: str,
) -> bytes:

    """
    Gera um Excel contendo somente as alterações.

    Estrutura:

    - Cabeçalho igual ao arquivo antigo.
    - Linhas originais permanecem nas mesmas posições.
    - Células inalteradas ficam vazias.
    - Alterações ficam amarelas.
    - Adições ficam verdes.
    - Remoções ficam vermelhas e exibem o valor antigo.
    - Linhas novas são adicionadas ao final em verde.
    """

    old_wb = load_workbook(
        BytesIO(old_bytes)
    )

    if sheet_name not in old_wb.sheetnames:
        raise ValueError(
            f"Aba '{sheet_name}' não encontrada "
            "no arquivo antigo."
        )

    old_ws = old_wb[sheet_name]

    header_row = _header_row(
        old_ws
    )

    old_headers = _headers(
        old_ws,
        header_row,
    )

    # ========================================================
    # NOVO WORKBOOK
    # ========================================================

    wb = Workbook()

    ws = wb.active
    ws.title = sheet_name

    # ========================================================
    # COPIAR CABEÇALHO
    # ========================================================

    max_col = (
        old_ws.max_column or 1
    )

    for col in range(
        1,
        max_col + 1,
    ):

        source = old_ws.cell(
            header_row,
            col,
        )

        target = ws.cell(
            1,
            col,
        )

        target.value = source.value

        _copy_cell_style(
            source,
            target,
        )

        source_letter = get_column_letter(
            col
        )

        target_letter = get_column_letter(
            col
        )

        source_dimension = (
            old_ws.column_dimensions[
                source_letter
            ]
        )

        target_dimension = (
            ws.column_dimensions[
                target_letter
            ]
        )

        if source_dimension.width:
            target_dimension.width = (
                source_dimension.width
            )

        target_dimension.hidden = (
            source_dimension.hidden is True
        )

    # Altura do cabeçalho.
    if old_ws.row_dimensions[
        header_row
    ].height:

        ws.row_dimensions[1].height = (
            old_ws.row_dimensions[
                header_row
            ].height
        )

    # ========================================================
    # MAPEAMENTO DE COLUNAS
    # ========================================================

    headers = {}

    for col in range(
        1,
        max_col + 1,
    ):

        value = ws.cell(
            1,
            col,
        ).value

        if value is None:
            continue

        normalized = (
            _normalize_column_name(
                value
            )
        )

        if normalized not in headers:
            headers[normalized] = col

    next_free_col = (
        max_col + 1
    )

    def ensure_col(name: str) -> int:

        nonlocal next_free_col

        normalized = (
            _normalize_column_name(
                name
            )
        )

        if normalized in headers:
            return headers[
                normalized
            ]

        col = next_free_col

        next_free_col += 1

        # Copia estilo da primeira coluna.
        if headers:

            reference_col = next(
                iter(headers.values())
            )

            _copy_header_style(
                ws,
                reference_col,
                col,
                header_row=1,
            )

        ws.cell(
            1,
            col,
        ).value = name

        headers[normalized] = col

        return col

    # ========================================================
    # CÉLULAS ALTERADAS
    # ========================================================

    for change in result.modified:

        target_row = (
            change.row + 2
        )

        col = ensure_col(
            change.column
        )

        _write_cell(
            ws,
            target_row,
            col,
            change.new,
            FILL_CHANGED,
        )

    # ========================================================
    # CÉLULAS ADICIONADAS
    # ========================================================

    for change in result.added_cells:

        target_row = (
            change.row + 2
        )

        col = ensure_col(
            change.column
        )

        _write_cell(
            ws,
            target_row,
            col,
            change.new,
            FILL_ADDED,
        )

    # ========================================================
    # CÉLULAS REMOVIDAS
    # ========================================================

    for change in result.removed_cells:

        target_row = (
            change.row + 2
        )

        col = ensure_col(
            change.column
        )

        # Como o novo valor é vazio,
        # mostramos o valor antigo para
        # que a remoção seja visualmente identificável.
        _write_cell(
            ws,
            target_row,
            col,
            change.old,
            FILL_REMOVED,
        )

    # ========================================================
    # LINHAS ADICIONADAS
    # ========================================================

    added = result.added_rows.copy()

    added = added.drop(
        columns=["_linha_excel"],
        errors="ignore",
    )

    # As linhas antigas ocupam:
    #
    # linha 1 = cabeçalho
    # linha 2 em diante = dados antigos
    #
    # Portanto as linhas novas começam depois
    # da quantidade de linhas antigas.

    old_data_rows = 0

    if not added.empty or result.modified or result.added_cells:

        old_data_rows = (
            max(
                len(
                    result.removed_rows
                ),
                0,
            )
        )

    # O melhor ponto de partida é descobrir
    # quantas linhas existiam originalmente
    # no arquivo antigo.
    last_old_row = _last_data_row(
        old_ws,
        header_row,
    )

    original_data_count = max(
        last_old_row - header_row,
        0,
    )

    next_added_row = (
        original_data_count + 2
    )

    for _, row in added.iterrows():

        for col_name, value in row.items():

            col = ensure_col(
                str(col_name)
            )

            _write_cell(
                ws,
                next_added_row,
                col,
                _cell_str(value),
                FILL_ADDED,
            )

        next_added_row += 1

    # ========================================================
    # SALVAR
    # ========================================================

    out = BytesIO()

    wb.save(out)

    return out.getvalue()
